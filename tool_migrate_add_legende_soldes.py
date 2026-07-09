#!/usr/bin/env python3-uno
"""Migration idempotente : ajoute 3 lignes de légende des libellés #Solde dans la
table CONV (feuille Patrimoine, cols J-L), juste sous « Fond jaune – signalement ».

Documente pour l'utilisateur les libellés de la colonne #Solde de la feuille
Opérations (#141) :
  - « Relevé compte »   : solde lu d'un relevé de compte
  - « Σ Solde calculé » : solde calculé à l'import, aucun relevé attendu
  - « ⚠ Solde calculé » : solde calculé à l'import, relevé attendu mais absent

Idempotent : ne fait rien si « Σ Solde calculé » est déjà présent dans la table CONV.
Insertion SCOPÉE aux colonnes J-L (les colonnes A-H de Patrimoine restent intactes).
Les named ranges CONVnom/CONVcell/CONVlégende s'étendent automatiquement.

Usage :
    python3 tool_migrate_add_legende_soldes.py ~/Compta/comptes.xlsm [--dry-run]

--dry-run : sonde openpyxl SANS LibreOffice (rc 3 = changerait, 0 = déjà à jour).
"""
import argparse
import sys
from pathlib import Path

MARKER = 'Σ Solde calculé'
SHEET = 'Patrimoine'

# (CONVcell = libellé, CONVlégende = explication) — ordre d'affichage
LEGENDE = [
    ('Relevé compte',   'Solde de relevé de compte (feuille Opérations)'),
    ('Σ Solde calculé', "Solde calculé à l'import quand aucun relevé n'est attendu"),
    ('⚠ Solde calculé', "Solde calculé à l'import quand le relevé attendu est absent"),
]
BLUE = 0x0432FF   # texte bleu des libellés (colonne CONVcell)
WHITE = 0xFFFFFF  # fond blanc


def _already_present(xlsm_path):
    """Sonde openpyxl (SANS LibreOffice) : le MARKER est-il déjà dans CONVcell ?
    Retourne True si présent (migration inutile), False sinon."""
    import openpyxl
    from inc_excel_schema import ColResolver
    wb = openpyxl.load_workbook(xlsm_path, data_only=True)
    try:
        ws = wb[SHEET]
        cr = ColResolver.from_openpyxl(wb)
        col = cr.col('CONVcell')            # 1-indexed
        r0, r1 = cr.rows('CONVcell')
        if r0 is None:
            return True                      # pas de table CONV → rien à faire
        for r in range(r0, r1 + 1):
            if str(ws.cell(r, col).value or '').strip() == MARKER:
                return True
        return False
    finally:
        wb.close()


def _apply(xlsm_path):
    """Insère les 3 lignes (contenu + format) via UNO. Idempotent."""
    from inc_uno import UnoDocument, check_lock_file
    from inc_excel_schema import ColResolver, uno_row

    busy = check_lock_file(xlsm_path)
    if busy:
        print(f"✗ Classeur ouvert ({busy}) — ferme LibreOffice puis relance.", file=sys.stderr)
        return 1

    with UnoDocument(str(xlsm_path)) as doc:
        from com.sun.star.sheet.CellInsertMode import DOWN
        from com.sun.star.table.CellHoriJustify import CENTER

        xdoc = doc.document
        cr = ColResolver.from_uno(xdoc)          # 0-indexed
        ws = doc.get_sheet(SHEET)
        col_cell = cr.col('CONVcell')            # K (0-idx)
        col_leg = cr.col('CONVlégende')          # L (0-idx)
        letJ = cr.letter('CONVnom')              # J
        letL = cr.letter('CONVlégende')          # L
        r0, r1 = cr.rows('CONVcell')             # 1-indexed

        # Idempotence : déjà présent ?
        for r in range(r0, r1 + 1):
            if ws.getCellByPosition(col_cell, uno_row(r)).getString().strip() == MARKER:
                print("✓ Légende déjà présente — rien à faire.")
                return 0

        # Ancre : la ligne « … signalement » (fond jaune), sinon avant « … ne pas renommer »
        anchor = None
        for r in range(r0, r1 + 1):
            if 'signalement' in ws.getCellByPosition(col_leg, uno_row(r)).getString().lower():
                anchor = r
                break
        if anchor is None:
            for r in range(r0, r1 + 1):
                if 'ne pas renommer' in ws.getCellByPosition(col_cell, uno_row(r)).getString().lower():
                    anchor = r - 1
                    break
        if anchor is None:
            print("✗ Point d'ancrage introuvable dans la table CONV.", file=sys.stderr)
            return 1

        at = anchor + 1                           # 1-indexed : 1re ligne insérée
        n = len(LEGENDE)

        # Insertion scopée J:L (A-H intactes) — pousse les cellules J-L vers le bas
        rng = ws.getCellRangeByName(f"{letJ}{at}:{letL}{at + n - 1}")
        ws.insertCells(rng.RangeAddress, DOWN)

        for i, (label, leg) in enumerate(LEGENDE):
            rr = uno_row(at + i)
            kc = ws.getCellByPosition(col_cell, rr)
            lc = ws.getCellByPosition(col_leg, rr)
            kc.setString(label)
            lc.setString(leg)
            for c in (kc, lc):
                c.CellBackColor = WHITE
                c.CharFontName = 'Arial'
                c.CharHeight = 11
            kc.CharColor = BLUE
            kc.HoriJustify = CENTER

        doc.save()
        print(f"✓ Légende ajoutée ({n} lignes) sous « signalement » (L{at}..L{at + n - 1}).")

    # Hors du `with` (fichier fermé) : recadrer la vue salie par le save UNO
    # (patch ZIP-XML pur, idempotent). Cf. doctrine « --reframe avant livraison ».
    from tool_fix_formats import frame_views
    frame_views(str(xlsm_path), verbose=False)
    return 0


def main():
    ap = argparse.ArgumentParser(description="Ajoute la légende des libellés #Solde (table CONV).")
    ap.add_argument('xlsm', help='chemin du classeur comptes.xlsm')
    ap.add_argument('--dry-run', action='store_true',
                    help='sonde openpyxl (rc 3 = changerait, 0 = déjà à jour), sans LibreOffice')
    args = ap.parse_args()

    p = Path(args.xlsm).expanduser()
    if not p.exists():
        print(f"✗ Introuvable : {p}", file=sys.stderr)
        return 1

    if args.dry_run:
        return 0 if _already_present(p) else 3

    return _apply(p)


if __name__ == '__main__':
    sys.exit(main())
