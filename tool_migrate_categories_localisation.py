#!/usr/bin/env python3-uno
"""Migration idempotente (#208) : localiser les erreurs de saisie DEPUIS le classeur.

Deux volets, une même définition (les prédicats des verdicts) :

1. Feuille Contrôles — le contrôle CATÉGORIES devient un AGRÉGATEUR à
   sous-lignes, comme DIVERS et BALANCES :
     CATÉGORIES                         K = ✗ si L > 0 ; L = nb de sous-lignes en alarme
         Manquantes                     L = opérations (date numérique) sans catégorie
         Inconnues                      L = catégories hors CATnom (méta `#…` exclues)
         Écart Budget € (année glissante)  L/M = le calcul historique, conservé verbatim
   Avant : la position 2 de A1 ne pesait que des MONTANTS sur 365 j (Budget) —
   une catégorie inconnue de montant nul, compensé ou vieille d'un an passait,
   et rien ne disait QUELLE ligne fautait. Les deux nouveaux comptages sont
   cellule-exacts (modèle : INCONNUS L76) → verdict et localisation concordent.

   INCONNUS (Comptes, Devises) compte aussi les comptes VIDES d'une opération
   réelle (date numérique, hors méta-opérations `#…` comme #Balance) et les
   DEVISES vides ou absentes de COTcode : une ligne à date et montant sans
   compte, ou une devise mal saisie, n'était comptée par aucun contrôle
   (l'alarme Cotations ne lit que Plus_value et Avoirs).

2. Feuille Opérations — sur les plages nommées ENTIÈRES (Date, Réf., Catégorie,
   Compte) : un format conditionnel (rouge = la cellule que le verdict compte)
   et une validation par formule « Arrêter » (Catégorie ∈ CATnom ou `#…`,
   Compte ∈ CTRL1compte, Date bornée). Définition unique et écrivain :
   `inc_formats.OP_CELL_RULES` / `apply_operations_cell_rules` — que l'import
   re-joue à chaque run (LibreOffice tronque les validations à sa sauvegarde).
   Les miettes héritées (Réf. `-` en 7 morceaux, « Hors compte » mort) sont
   remplacées.

Localisation — DOCTRINE NAMED RANGES : le bloc CTRL2 se déplace verticalement
(CTRL1 grossit avec le nombre de comptes) → la ligne CATÉGORIES est trouvée
dans `CTRL2type` (col J), jamais par une cellule en dur. Les NR CTRL2* et la
synthèse (K$81 → recalée par LO) s'étendent seuls à l'insertion.

Ordre d'écriture : openpyxl D'ABORD (volet 2), UNO ENSUITE (volet 1 + recalcul +
sauvegarde) — une sauvegarde openpyxl n'écrit pas les valeurs en cache des
formules ; la sauvegarde UNO qui suit les restaure.

Idempotent : volet 1 sauté si la sous-ligne « Manquantes » suit déjà CATÉGORIES ;
volet 2 sauté si les règles sont conformes.

Usage :
    python3 tool_migrate_categories_localisation.py ~/Compta/comptes.xlsm [--dry-run]

--dry-run : sonde openpyxl SANS LibreOffice (rc 3 = changerait, 0 = déjà à jour).
xlsx (classeur exemple) et xlsm.
"""
import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

SHEET = 'Contrôles'
HEAD = 'CATÉGORIES'               # préfixe du libellé de tête (col J, NR CTRL2type)
HEAD_LABEL = 'CATÉGORIES'         # libellé cible (le suffixe « (année glissante) » descend en sous-ligne)
INDENT = '    '                   # 4 espaces : sous-ligne (pattern DIVERS/BALANCES)
SUB_MANQ = 'Manquantes'
SUB_INC = 'Inconnues'
SUB_ECART = 'Écart Budget € (année glissante)'
HEAD_INC = 'INCONNUS'             # tête INCONNUS (Comptes) : compte aussi les comptes VIDES
F_INC_VIDES = '+SUMPRODUCT(ISNUMBER(OPdate)*(OPcompte="")*(LEFT(OPcatégorie;1)<>"#"))'
F_INC_VIDES_V0 = '+SUMPRODUCT(ISNUMBER(OPdate)*(OPcompte=""))'   # 1er jet (jamais livré) → remplacé
F_INC_DEVISE = ('+SUMPRODUCT((COUNTIF(COTcode;OPdevise)=0)*(OPdevise<>""))'
                '+SUMPRODUCT(ISNUMBER(OPdate)*(OPdevise="")*(LEFT(OPcatégorie;1)<>"#"))')
HEAD_INC_LABEL = 'INCONNUS (Comptes, Devises)'
LABEL_NR = 'CTRL2type'
DISPL_NR = 'CTRL2affichage'
VALUE_NR = 'CTRL2general'
DRILL_NR = 'CTRL2drill'           # colonne M (montant par devise, sélecteur M$62)

# Formules des deux comptages (séparateur UNO `;`). Miroir des prédicats
# d'inc_formats.OP_CELL_RULES (catégorie) : opération réelle = date numérique ;
# méta-catégories `#…` (plage Spéciale) exclues, comme dans le calcul Budget.
F_MANQ = '=SUMPRODUCT(ISNUMBER(OPdate)*(OPcatégorie=""))'
F_INC = ('=SUMPRODUCT((COUNTIF(CATnom;OPcatégorie)=0)*(OPcatégorie<>"")'
         '*(LEFT(OPcatégorie;1)<>"#"))')


def _nr_bounds_openpyxl(wb, name):
    """(col_1indexed, start_row, end_row) d'un named range colonne, ou None."""
    import re
    from openpyxl.utils import column_index_from_string
    dn = wb.defined_names.get(name)
    if dn is None:
        return None
    m = re.search(r"\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)", str(dn.value))
    if not m:
        return None
    return column_index_from_string(m.group(1)), int(m.group(2)), int(m.group(4))


def _probe(path):
    """Sonde openpyxl (SANS LibreOffice). Retourne (controles_a_faire, regles_a_poser)."""
    import openpyxl
    from inc_formats import apply_operations_cell_rules
    wb = openpyxl.load_workbook(path, keep_vba=str(path).endswith('.xlsm'))
    try:
        ctrl = False
        if SHEET in wb.sheetnames:
            ws = wb[SHEET]
            lab = _nr_bounds_openpyxl(wb, LABEL_NR)
            if lab:
                lcol, start, end = lab
                val = _nr_bounds_openpyxl(wb, VALUE_NR)
                for r in range(start, end + 1):
                    v = str(ws.cell(r, lcol).value or '').strip()
                    if v.startswith(HEAD):
                        nxt = str(ws.cell(r + 1, lcol).value or '').strip()
                        ctrl = ctrl or (nxt != SUB_MANQ)
                    elif v.startswith(HEAD_INC) and val:
                        f = str(ws.cell(r, val[0]).value or '')
                        ctrl = ctrl or (F_INC_VIDES.replace(';', ',') not in f) \
                                    or (F_INC_DEVISE.replace(';', ',') not in f) \
                                    or v != HEAD_INC_LABEL
        cells = bool(apply_operations_cell_rules(wb, apply=False))
        return ctrl, cells
    finally:
        wb.close()


def _apply_cells(path):
    """Volet 2 — openpyxl : pose des règles cellule d'Opérations."""
    import openpyxl
    from inc_formats import apply_operations_cell_rules
    is_xlsm = str(path).endswith('.xlsm')
    wb = openpyxl.load_workbook(path, keep_vba=is_xlsm)
    changes = apply_operations_cell_rules(wb, apply=True)
    if changes:
        wb.save(path)
    wb.close()
    return changes


def _apply_inconnus(ws, jc, lc, start, end):
    """Volet 1 bis — INCONNUS (Comptes) compte aussi les comptes vides. Idempotent."""
    for r1 in range(start, end + 1):
        if ws.getCellByPosition(jc, r1 - 1).getString().strip().startswith(HEAD_INC):
            cell = ws.getCellByPosition(lc, r1 - 1)
            f = cell.getFormula()
            if not f.startswith('='):
                raise RuntimeError(f"L{r1} (INCONNUS) sans formule ({f!r})")
            out = []
            if F_INC_VIDES not in f:
                f = f.replace(F_INC_VIDES_V0, F_INC_VIDES) if F_INC_VIDES_V0 in f else f + F_INC_VIDES
                out.append(f"~ {SHEET} L{r1} : INCONNUS compte aussi les comptes vides")
            if F_INC_DEVISE not in f:
                f += F_INC_DEVISE
                out.append(f"~ {SHEET} L{r1} : INCONNUS compte aussi les devises vides ou inconnues")
            if out:
                cell.setFormula(f)
            jcell = ws.getCellByPosition(jc, r1 - 1)
            if jcell.getString().strip() != HEAD_INC_LABEL:
                old = jcell.getString().strip()
                jcell.setString(HEAD_INC_LABEL)
                out.append(f"~ {SHEET} J{r1} : « {old} » → « {HEAD_INC_LABEL} »")
            return out
    raise RuntimeError(f"libellé « {HEAD_INC} » introuvable dans {LABEL_NR}")


def _apply_controles(doc):
    """Volet 1 — UNO : CATÉGORIES en agrégateur à 3 sous-lignes + INCONNUS étendu. Retourne list[str]."""
    from inc_uno import get_col_range_bounds, copy_row_style
    xdoc = doc.document
    ws = doc.get_sheet(SHEET)
    lab_b = get_col_range_bounds(xdoc, LABEL_NR)
    dis_b = get_col_range_bounds(xdoc, DISPL_NR)
    val_b = get_col_range_bounds(xdoc, VALUE_NR)
    dri_b = get_col_range_bounds(xdoc, DRILL_NR)
    if not (lab_b and dis_b and val_b and dri_b):
        raise RuntimeError(f"named ranges {LABEL_NR}/{DISPL_NR}/{VALUE_NR}/{DRILL_NR} absents "
                           f"({SHEET}) — classeur trop ancien ?")
    jc, kc, lc, mc = lab_b[1], dis_b[1], val_b[1], dri_b[1]      # cols 0-indexées
    start, end = lab_b[2], lab_b[3]                              # lignes 1-indexées

    r_head = None
    for r1 in range(start, end + 1):
        if ws.getCellByPosition(jc, r1 - 1).getString().strip().startswith(HEAD):
            r_head = r1
            break
    if r_head is None:
        raise RuntimeError(f"libellé « {HEAD} » introuvable dans {LABEL_NR}")
    if ws.getCellByPosition(jc, r_head).getString().strip() == SUB_MANQ:
        return _apply_inconnus(ws, jc, lc, start, end)          # sous-lignes déjà en place

    def cell(c, r1):
        return ws.getCellByPosition(c, r1 - 1)

    # Repères AVANT insertion (0-indexées pour copy_row_style)
    r0_head = r_head - 1
    old_l = cell(lc, r_head).getFormula()          # écart Budget (verbatim, locale UNO)
    old_m = cell(mc, r_head).getFormula()          # écart par devise (verbatim)
    if not old_l.startswith('='):
        raise RuntimeError(f"L{r_head} (CATÉGORIES) sans formule ({old_l!r})")
    fmt_l_ecart = cell(lc, r_head).NumberFormat
    fmt_m_ecart = cell(mc, r_head).NumberFormat
    cf_m_ecart = cell(mc, r_head).ConditionalFormat   # formats par devise ($M$62="EUR"…)

    # Modèles de style : une sous-ligne existante (« Date hors période », 1re de
    # DIVERS) et une tête à compteur entier (DIVERS). Cherchés par libellé.
    r_sub_model = r_head_model = None
    for r1 in range(start, end + 1):
        v = cell(jc, r1).getString()
        if v.strip() == 'Date hors période' and r_sub_model is None:
            r_sub_model = r1
        if v.strip() == 'DIVERS' and r_head_model is None:
            r_head_model = r1

    # Insertion de 3 lignes juste APRÈS la tête (index 0-based = r_head)
    ws.Rows.insertByIndex(r_head, 3)
    r_manq, r_inc, r_ecart = r_head + 1, r_head + 2, r_head + 3
    # Les modèles situés sous la tête ont glissé de 3
    if r_sub_model and r_sub_model > r_head:
        r_sub_model += 3
    if r_head_model and r_head_model > r_head:
        r_head_model += 3

    changes = []
    # Style des sous-lignes (J..M) depuis la sous-ligne modèle
    if r_sub_model:
        for r1 in (r_manq, r_inc, r_ecart):
            copy_row_style(ws, r_sub_model - 1, r1 - 1, col_start=jc, col_end=mc + 1)
            cf_src = cell(lc, r_sub_model).ConditionalFormat      # « ≠ 0 » en rouge
            if cf_src.Count:
                cell(lc, r1).ConditionalFormat = cf_src
    else:
        changes.append("⚠ sous-ligne modèle « Date hors période » introuvable : style non copié")

    # Sous-lignes : J indenté, K vide, L = comptage / calcul historique
    cell(jc, r_manq).setString(INDENT + SUB_MANQ)
    cell(kc, r_manq).setString('')
    cell(lc, r_manq).setFormula(F_MANQ)
    cell(jc, r_inc).setString(INDENT + SUB_INC)
    cell(kc, r_inc).setString('')
    cell(lc, r_inc).setFormula(F_INC)
    cell(jc, r_ecart).setString(INDENT + SUB_ECART)
    cell(kc, r_ecart).setString('')
    cell(lc, r_ecart).setFormula(old_l)
    cell(lc, r_ecart).NumberFormat = fmt_l_ecart
    if old_m.startswith('='):
        cell(mc, r_ecart).setFormula(old_m)
        cell(mc, r_ecart).NumberFormat = fmt_m_ecart
        if cf_m_ecart.Count:
            cell(mc, r_ecart).ConditionalFormat = cf_m_ecart
    changes += [f"+ {SHEET} L{r_manq} : sous-ligne « {SUB_MANQ} »",
                f"+ {SHEET} L{r_inc} : sous-ligne « {SUB_INC} »",
                f"+ {SHEET} L{r_ecart}/M{r_ecart} : sous-ligne « {SUB_ECART} » (calcul historique déplacé)"]

    # Tête : libellé court, K agrégateur, L = nb de sous-lignes en alarme, M vidée
    old_label = cell(jc, r_head).getString().strip()
    if old_label != HEAD_LABEL:
        cell(jc, r_head).setString(HEAD_LABEL)
        changes.append(f"~ {SHEET} J{r_head} : « {old_label} » → « {HEAD_LABEL} »")
    cell(kc, r_head).setFormula(f'=IF(L{r_head}>0;"✗";"✓")')
    cell(lc, r_head).setFormula(f'=IF(L{r_manq}>0;1;0)+IF(L{r_inc}>0;1;0)'
                                f'+IF(ABS(L{r_ecart})>=1;1;0)')
    if r_head_model:
        cell(lc, r_head).NumberFormat = cell(lc, r_head_model).NumberFormat
        cell(mc, r_head).NumberFormat = cell(mc, r_head_model).NumberFormat
    cell(mc, r_head).setString('')
    cf_head_m = cell(mc, r_head).ConditionalFormat
    if cf_head_m.Count:
        cf_head_m.clear()
        cell(mc, r_head).ConditionalFormat = cf_head_m
    changes.append(f"~ {SHEET} K{r_head}/L{r_head} : CATÉGORIES en agrégateur "
                   f"(✗ si une sous-ligne est en alarme)")
    _restrict_head_cf(ws, jc, kc, r_head, changes)
    changes += _apply_inconnus(ws, jc, lc, start, end + 3)
    return changes


def _restrict_head_cf(ws, jc, kc, r_head, changes):
    """LibreOffice ÉTEND la CF d'alarme de K{tête} aux lignes insérées juste
    dessous (K65 → K65:K68). Or `FIND("✗"|"⚠"; INDIRECT("RC";0))` sur une plage
    multi-lignes se déclenche par effet de bord sur les sous-lignes voisines
    quand la tête s'allume (vécu v4.1.0, section 7). → fissionner : la CF ne
    reste que sur la tête. Même approche que tool_migrate_v4.1.0."""
    import re
    from inc_uno import set_alarm_cf
    k_letter = ''
    n = kc + 1
    while n:
        n, rem = divmod(n - 1, 26)
        k_letter = chr(65 + rem) + k_letter
    cfs = ws.ConditionalFormats
    for cf in list(cfs.ConditionalFormats):
        addr = cf.Range.getRangeAddressesAsString()
        m = re.match(r"^[^.]+\.([A-Z]+)(\d+):([A-Z]+)(\d+)$", addr)
        if not m or m.group(1) != k_letter or m.group(3) != k_letter:
            continue
        row_s, row_e = int(m.group(2)), int(m.group(4))
        if row_s == row_e or not (row_s <= r_head <= row_e):
            continue
        heads = [r for r in range(row_s, row_e + 1)
                 if (v := ws.getCellByPosition(jc, r - 1).getString())
                 and not v.startswith(' ') and v.strip() != '⚓']
        if not heads or len(heads) == row_e - row_s + 1:
            continue
        cfs.removeByID(cf.ID)
        for r in heads:
            set_alarm_cf(ws.getCellByPosition(kc, r - 1))
        changes.append(f"~ {SHEET} CF {k_letter}{row_s}:{k_letter}{row_e} → restreinte à "
                       + ','.join(f'{k_letter}{r}' for r in heads))


def _apply(path):
    from inc_uno import UnoDocument, check_lock_file, require_libreoffice_min

    if check_lock_file(path):
        print("✗ Classeur ouvert (LibreOffice) — ferme-le puis relance.", file=sys.stderr)
        return 1
    require_libreoffice_min(24, 8)

    # Volet 2 d'abord (openpyxl), puis volet 1 (UNO, qui recalcule et sauve)
    cell_changes = _apply_cells(path)
    for c in cell_changes:
        print(f"  + Opérations : {c}")

    with UnoDocument(str(path)) as doc:
        try:
            ctrl_changes = _apply_controles(doc)
        except RuntimeError as e:
            print(f"✗ {e}", file=sys.stderr)
            return 1
        for c in ctrl_changes:
            print(f"  {c}")
        doc.calculate_all()
        doc.save()

    # Hors du `with` : recadrer la vue salie par le save UNO (patch ZIP-XML pur)
    from tool_fix_formats import frame_views
    frame_views(str(path), verbose=False)

    if not cell_changes and not ctrl_changes:
        print("✓ Déjà à jour — rien à faire.")
    else:
        print(f"✓ Migration appliquée ({len(ctrl_changes)} changement(s) Contrôles, "
              f"{len(cell_changes)} règle(s) Opérations).")
    return 0


def main():
    ap = argparse.ArgumentParser(
        description="CATÉGORIES en agrégateur à sous-lignes + localisation des erreurs "
                    "de saisie dans Opérations (formats conditionnels, listes) — #208.")
    ap.add_argument('xlsm', help='chemin du classeur (comptes.xlsm, ou .xlsx exemple)')
    ap.add_argument('--dry-run', action='store_true',
                    help='sonde openpyxl (rc 3 = changerait, 0 = déjà à jour), sans LibreOffice')
    args = ap.parse_args()

    p = Path(args.xlsm).expanduser()
    if not p.exists():
        print(f"✗ Introuvable : {p}", file=sys.stderr)
        return 1

    if args.dry_run:
        ctrl, cells = _probe(p)
        return 3 if (ctrl or cells) else 0

    return _apply(p)


if __name__ == '__main__':
    sys.exit(main())
